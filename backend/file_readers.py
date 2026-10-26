from pathlib import Path
from pptx import Presentation
import pypdfium2 as pdfium
from pypdf import PdfReader as PypdfReader
import numpy as np
from llama_index.core import Document
from llama_index.core.readers.base import BaseReader
from llama_index.core import SimpleDirectoryReader


SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".pptx", ".jpeg", ".jpg", ".png"}


#for .pptx files
class SlideReader(BaseReader):
    """One Document per slide: slide text, table rows, and speaker notes."""

    def load_data(self, file: Path, extra_info: dict | None = None, **kwargs) -> list[Document]:
        prs = Presentation(str(file))
        docs = []
        for num, slide in enumerate(prs.slides, start=1):
            lines = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        lines.append(text)
                elif shape.has_table:
                    for row in shape.table.rows:
                        lines.append(" | ".join(cell.text.strip() for cell in row.cells))

            # get the speaker notes
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    lines.append(f"Speaker notes: {notes}")

            if lines:   # skip slides that are only images
                docs.append(Document(
                    text="\n".join(lines),
                    metadata={**(extra_info or {}), "page_label": str(num)},
                ))
        return docs


_ocr = None

def get_ocr():
    """Load easyocr once, on first use, and share it between every reader."""
    global _ocr   # reassign the module variable instead of creating a local one
    if _ocr is None:
        import easyocr   # heavy import (pulls in torch), so only when actually needed
        _ocr = easyocr.Reader(["en"], gpu=False)
    return _ocr

class ImageReader(BaseReader):

    def load_data(self, file: Path, extra_info: dict | None = None, **kwargs) -> list[Document]:
        lines = get_ocr().readtext(str(file), detail=0, paragraph=True)
        text = "\n".join(lines).strip()
        if not text:
            return []
        return [Document(text=text, metadata=extra_info or {})]

class OcrPdfReader(BaseReader):
    """One Document per page: the text layer when there is one, OCR when there isn't."""

    def load_data(self, file: Path, extra_info:dict | None = None, **kwargs)->list[Document]:
        
        #go through each page of the pdf
        reader = PypdfReader(str(file))
        pdf = pdfium.PdfDocument(str(file)) #for ocr scanning

        docs = []

        try:
            i = 0
            while i < len(reader.pages):
                #try to extract text with pypdf
                text = (reader.pages[i].extract_text() or "").strip()  # "or ''" in case it returns None
                method = "text"

                #if it can't, could be a scanned pdf -> try converting to PIL image for easyocr to read from
                if not text:

                     #render it big so the ocr can see easier and convert to a PIL Image
                    image = pdf[i].render(scale=2).to_pil()  
                    lines = get_ocr().readtext(np.array(image), detail=0, paragraph=True)
                    text = "\n".join(lines).strip()   # normalize to a string immediately
                    method = "ocr"

                        
                if text: 
                    docs.append(Document(
                        text=text,
                        metadata={**(extra_info or {}),
                                  "page_label": str(i + 1),
                                  "extraction": method},
                    ))
                i += 1
        # so the file closes even if an issue happens
        finally:                
            pdf.close() 

        return docs

# The readers override the defaults for these extensions; everything else
_image_reader = ImageReader() 
# .docx still uses the built-in reader
# but assigns  Slidereader to .pptx format
FILE_EXTRACTORS = {".pdf" : OcrPdfReader(), ".pptx": SlideReader(), ".jpeg": _image_reader, ".jpg": _image_reader, ".png": _image_reader} 


def load_documents(path: str) -> list[Document]:
    docs = SimpleDirectoryReader(
        input_files=[path],
        file_extractor=FILE_EXTRACTORS,
    ).load_data()
   #safety net for readers in case of empty content
    return [d for d in docs if d.text.strip()]